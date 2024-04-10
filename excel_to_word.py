import pandas as pd
import docx

# Load the Excel file
df = pd.read_excel("C:\\Users\\2113196\\Downloads\\intestinal_performation_case_line_listing_table_cumulative.xlsx")

# Create a new Word document or open an existing one
doc = docx.Document()

# Add a table to the document and create a reference variable
# Extra row is for the header row
t = doc.add_table(df.shape[0]+1, df.shape[1])
t.style = 'Table Grid'

# Add the header rows
for j in range(df.shape[-1]):
    t.cell(0, j).text = str(df.columns[j])

# Add the rest of the DataFrame
for i in range(df.shape[0]):
    for j in range(df.shape[-1]):
        t.cell(i+1, j).text = str(df.values[i, j])

# Save the document
doc.save("C:\\Users\\2113196\\Downloads\\aaaaaaaaaaa.docx")